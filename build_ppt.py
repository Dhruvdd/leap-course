from pptx import Presentation
from pptx.util import Pt, Inches
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
import re

from pptx import Presentation
from pptx.util import Pt, Inches
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
import re

def parse_slides(slides_text):
    blocks = re.split(r"(?=^Slide\s*\d+\s*[-–])", slides_text, flags=re.MULTILINE)
    slides = []
    for block in blocks:
        lines = [l for l in block.strip().split('\n') if l.strip()]
        if not lines or not lines[0].startswith("Slide "):
            continue
        title_line = lines[0]
        match = re.match(r"Slide\s*\d+\s*[-–]\s*(.+)", title_line)
        title = match.group(1).strip() if match else title_line
        slide_content = "\n".join(lines[1:]).strip()
        slides.append({"title": title, "content": slide_content})
    return slides

def get_bullet_level(line):
    leading_spaces = len(line) - len(line.lstrip(' '))
    level = leading_spaces // 4
    clean = line.strip()
    if clean.startswith("• "):
        clean = clean[2:]
    return clean, level

def set_background_image(slide, prs, image_path):
    slide_width = prs.slide_width
    slide_height = prs.slide_height
    pic = slide.shapes.add_picture(image_path, 0, 0, width=slide_width, height=slide_height)
    slide.shapes._spTree.remove(pic._element)
    slide.shapes._spTree.insert(2, pic._element)  # Send to very back

def add_white_box(slide, left, top, width, height, transparency=0.0):
    shape = slide.shapes.add_shape(
        1,  # Rectangle shape
        left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    shape.fill.transparency = transparency
    shape.line.color.rgb = RGBColor(255, 255, 255)
    shape.shadow.inherit = False
    return shape

def build_ppt_from_slides(slides_text, background_image, output_pptx):
    SLIDE_WIDTH = Inches(13.33)
    SLIDE_HEIGHT = Inches(7.5)
    BOX_LEFT = Inches(0.4)
    BOX_TOP = Inches(1)
    BOX_WIDTH = Inches(12.5)
    BOX_HEIGHT = Inches(6)

    slides = parse_slides(slides_text)
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT

    for slide_info in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[5])  # Use 'Title Only'
        set_background_image(slide, prs, background_image)
        add_white_box(slide, BOX_LEFT, BOX_TOP, BOX_WIDTH, BOX_HEIGHT)

        # Add Title
        title_shape = slide.shapes.add_textbox(BOX_LEFT + Inches(0.4), BOX_TOP + Inches(0.1), BOX_WIDTH - Inches(0.8), Inches(0.7))
        title_tf = title_shape.text_frame
        p = title_tf.paragraphs[0]
        p.text = slide_info["title"]
        p.font.size = Pt(30)
        p.font.bold = True
        p.font.color.rgb = RGBColor(23, 32, 42)
        p.alignment = PP_ALIGN.LEFT

        # Add Content
        content_shape = slide.shapes.add_textbox(BOX_LEFT + Inches(0.4), BOX_TOP + Inches(0.9), BOX_WIDTH - Inches(0.8), BOX_HEIGHT - Inches(1.1))
        tf = content_shape.text_frame
        tf.word_wrap = True
        tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = Inches(0.1)
        for line in slide_info["content"].split('\n'):
            if line.strip():
                text, level = get_bullet_level(line)
                para = tf.add_paragraph()
                para.text = text
                para.level = level
                para.font.size = Pt(18)
                para.font.color.rgb = RGBColor(43, 43, 43)
                para.alignment = PP_ALIGN.LEFT
        # Remove the first (empty) paragraph auto-added
        if tf.paragraphs and not tf.paragraphs[0].text:
            tf._element.remove(tf.paragraphs[0]._element)

    prs.save(output_pptx)
    print(f"PPTX created: {output_pptx}")